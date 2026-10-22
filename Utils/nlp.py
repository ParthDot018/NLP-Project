import io
import uuid
import nltk
import numpy as np
from pptx import Presentation
from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

nltk.download('punkt', quiet=True)
nltk.download('punkt_tab', quiet=True)
nltk.download('stopwords', quiet=True)

class NLP_Methods:
    def __init__(self):
        self.session_data = {}
    
    def _extract_text_from_pptx(self, file_bytes: bytes) -> dict:
        file_stream = io.BytesIO(file_bytes)
        try:
            prs = Presentation(file_stream)
        except Exception:
            return {}

        pages_text = {}
        for i, slide in enumerate(prs.slides, 1):
            slide_content = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text.strip()
                    if text:
                        slide_content.append(" ".join(text.split()))
            pages_text[i] = " ".join(slide_content)
        return pages_text

    def _extract_text_from_pdf(self, file_bytes: bytes) -> dict:
        file_stream = io.BytesIO(file_bytes)
        try:
            reader = PdfReader(file_stream)
        except Exception:
            return {}

        pages_text = {}
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text()
            if text:
                pages_text[i] = " ".join(text.split())
            else:
                pages_text[i] = ""
        return pages_text

    def _generate_topic_title(self, text_block: str) -> str:
        if not text_block.strip():
            return "General Overview"
        vectorizer = TfidfVectorizer(stop_words='english', max_features=3)
        try:
            vectorizer.fit_transform([text_block])
            keywords = vectorizer.get_feature_names_out()
            return " & ".join([word.title() for word in keywords])
        except ValueError:
            return "General Topic"

    def _summarize_topic(self, text: str, num_sentences: int = 2) -> str:
        if not text.strip():
            return "No text available."

        sentences = nltk.tokenize.sent_tokenize(text)
        if len(sentences) <= num_sentences:
            return text

        vectorizer = TfidfVectorizer(stop_words='english')
        try:
            tfidf_matrix = vectorizer.fit_transform(sentences)
        except ValueError:
            return " ".join(sentences[:num_sentences])

        sentence_scores = np.array(tfidf_matrix.sum(axis=1)).flatten()
        top_sentence_indices = sentence_scores.argsort()[-num_sentences:]
        top_sentence_indices.sort()
        
        return " ".join([sentences[i] for i in top_sentence_indices])

    def _group_into_topics(self, pages_text: dict, threshold: float = 0.15) -> list:
        vectorizer = TfidfVectorizer(stop_words='english')
        topics = []
        current_topic_id = str(uuid.uuid4())
        current_pages = []
        current_text = ""
        start_page = 1

        for page_num, text in pages_text.items():
            if not text.strip():
                current_pages.append(page_num)
                continue

            if not current_text:
                current_text = text
                current_pages.append(page_num)
                start_page = page_num
                continue
                
            try:
                tfidf_matrix = vectorizer.fit_transform([current_text, text])
                similarity = cosine_similarity(tfidf_matrix[0:1], tfidf_matrix[1:2])[0][0]
            except ValueError:
                similarity = 0.0

            if similarity >= threshold:
                current_text += " " + text
                current_pages.append(page_num)
            else:
                end_page = current_pages[-1]
                topics.append({
                    "id": current_topic_id,
                    "title": self._generate_topic_title(current_text),
                    "pages": f"{start_page}-{end_page}",
                    "raw_text": current_text,
                    "summary": self._summarize_topic(current_text) 
                })
                
                current_topic_id = str(uuid.uuid4())
                current_text = text
                current_pages = [page_num]
                start_page = page_num

        if current_text:
            topics.append({
                "id": current_topic_id,
                "title": self._generate_topic_title(current_text),
                "pages": f"{start_page}-{current_pages[-1]}",
                "raw_text": current_text,
                "summary": self._summarize_topic(current_text)
            })

        return topics

    
    def process_presentation(self, file_bytes: bytes, filename: str) -> list:
        if filename.lower().endswith('.pdf'):
            pages_text = self._extract_text_from_pdf(file_bytes)
        else:
            pages_text = self._extract_text_from_pptx(file_bytes)
            
        if not pages_text:
            return None

        topics = self._group_into_topics(pages_text)

        for t in topics:
            self.session_data[t["id"]] = t

        return topics
    
    def check_topic_exists(self, topic_id: str) -> bool:
        return topic_id in self.session_data

    def get_topic_text(self, topic_id: str) -> str:
        return self.session_data.get(topic_id, {}).get("raw_text", "")

    def get_extractive_answer(self, question: str, topic_text: str) -> str:
        sentences = nltk.tokenize.sent_tokenize(topic_text)
        if not sentences:
            return "No text available to answer the question."

        sentences.append(question)
        vectorizer = TfidfVectorizer(stop_words='english')
        
        try:
            tfidf_matrix = vectorizer.fit_transform(sentences)
        except ValueError:
            return "Not enough meaningful text to search."

        similarities = cosine_similarity(tfidf_matrix[-1], tfidf_matrix[:-1]).flatten()
        best_match_index = similarities.argmax()
        highest_score = similarities[best_match_index]
        sentences.pop()

        if highest_score > 0.1:
            return sentences[best_match_index]
        else:
            return "I couldn't find a direct answer to that in this topic's slides."